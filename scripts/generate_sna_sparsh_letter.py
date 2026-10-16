"""
SNA SPARSH Portal – SSO ID Mapping Request Letter
Sender : Executive Engineer, PWD District Division-II, Udaipur
Officer : Anil Khichi  |  Office ID : 7542
SSO ID  : RJRA201231041638
Email   : EEDN2PWDUDAIPUR@GMAIL.COM
"""

import os, shutil, tempfile
import sys; sys.path.insert(0, os.path.dirname(__file__))
from output_utils import publish_output
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR   = r"E:\Rajkumar\Legal-Correspondence_Manager\docs"
DOCX_PATH = os.path.join(OUT_DIR, "SNA_SPARSH_SSO_Mapping_DN2_Udaipur.docx")
PDF_PATH  = os.path.join(OUT_DIR, "SNA_SPARSH_SSO_Mapping_DN2_Udaipur.pdf")
FONT      = "Nirmala UI"

# ── Helpers ───────────────────────────────────────────────────────────────────
def apply_font(run, size, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT
    rPr    = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(slot), FONT)
    run.font.color.rgb = RGBColor(0, 0, 0)

def para(doc, text, bold=False, size=12,
         align=WD_ALIGN_PARAGRAPH.LEFT, sa=6, sb=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after  = Pt(sa)
    p.paragraph_format.space_before = Pt(sb)
    r = p.add_run(text)
    apply_font(r, size, bold)
    return p

def mixed(doc, parts, size=12,
          align=WD_ALIGN_PARAGRAPH.LEFT, sa=6, sb=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after  = Pt(sa)
    p.paragraph_format.space_before = Pt(sb)
    for text, bold in parts:
        r = p.add_run(text)
        apply_font(r, size, bold)
    return p

def kv(doc, key, val, sa=3):
    mixed(doc, [(key, False), (val, True)], size=12, sa=sa)

# ── Build Document ────────────────────────────────────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.page_width    = Cm(21.0)
sec.page_height   = Cm(29.7)
sec.left_margin   = Cm(3.0)
sec.right_margin  = Cm(2.5)
sec.top_margin    = Cm(2.5)
sec.bottom_margin = Cm(2.5)

sty = doc.styles["Normal"]
sty.font.name = FONT
sty.font.size = Pt(12)
sty._element.rPr.rFonts.set(qn("w:cs"), FONT)

# ── Office Header ─────────────────────────────────────────────────────────────
para(doc,
     "कार्यालय अधिशासी अभियन्ता सा.नि.वि. जिला खण्ड द्वितीय (उदयपुर)",
     bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, sa=2)

# Underline rule via paragraph bottom border
p_rule = doc.add_paragraph()
p_rule.paragraph_format.space_after  = Pt(8)
p_rule.paragraph_format.space_before = Pt(0)
pPr  = p_rule._p.get_or_add_pPr()
pBdr = OxmlElement("w:pBdr")
bot  = OxmlElement("w:bottom")
bot.set(qn("w:val"),   "single")
bot.set(qn("w:sz"),    "6")
bot.set(qn("w:space"), "1")
bot.set(qn("w:color"), "000000")
pBdr.append(bot)
pPr.append(pBdr)

# Kramank + Date (borderless 2-cell table)
tbl = doc.add_table(rows=1, cols=2)
tbl.style = "Table Grid"
tblPr    = tbl._tbl.tblPr
tblBdr   = OxmlElement("w:tblBorders")
for edge in ("top","left","bottom","right","insideH","insideV"):
    t = OxmlElement(f"w:{edge}")
    t.set(qn("w:val"), "none")
    tblBdr.append(t)
tblPr.append(tblBdr)

lc, rc = tbl.rows[0].cells[0], tbl.rows[0].cells[1]
lc.width, rc.width = Cm(9), Cm(6.5)
lp = lc.paragraphs[0]; lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
lr = lp.add_run("क्रमांक:– ___________"); apply_font(lr, 12)
rp = rc.paragraphs[0]; rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rr = rp.add_run("दिनांक:– 01 अक्टूबर 2026"); apply_font(rr, 12)

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ── Addressee ─────────────────────────────────────────────────────────────────
para(doc, "श्रीमान् मुख्य अभियन्ता (भवन)", size=12, sa=1)
para(doc, "सा.नि.वि. राजस्थान, जयपुर।",   size=12, sa=12)

# ── Subject ───────────────────────────────────────────────────────────────────
mixed(doc, [
    ("विषय:– ", True),
    ("SNA SPARSH PORTAL पर अधोहस्ताक्षरकर्ता की SSO ID मैपिंग करने के सम्बन्ध में।", False),
], size=12, sa=10)

# ── Salutation ────────────────────────────────────────────────────────────────
para(doc, "महोदय,", size=12, sa=8)

# ── Body ──────────────────────────────────────────────────────────────────────
mixed(doc, [
    ("        उपरोक्त विषयान्तर्गत निवेदन है, कि अधोहस्ताक्षरकर्ता विगत दो वर्ष से "
     "अधिशासी अभियन्ता, सा.नि.वि. जिला खण्ड द्वितीय, उदयपुर के पद पर कार्यरत हूँ। "
     "जिला खण्ड द्वितीय उदयपुर (", False),
    ("7542", True),
    (") के SNA SPARSH PORTAL पर संवेदकों के बिलों का भुगतान कार्य हेतु "
     "अधोहस्ताक्षरकर्ता की SSO ID :- ", False),
    ("RJRA201231041638", True),
    (" SNA SPARSH PORTAL पर मैप करवाने का कष्ट करावें।", False),
], size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, sa=14)

# ── Details block ─────────────────────────────────────────────────────────────
kv(doc, "Name:-  Sh. ",       "Anil Khichi")
kv(doc, "Name of Office:-  ", "Executive Engineer, PWD District Division-II, Udaipur")
kv(doc, "OFFICE ID :-  ",     "7542")
kv(doc, "SSO ID :-  ",        "RJRA201231041638")
kv(doc, "Email ID :-  ",      "EEDN2PWDUDAIPUR@GMAIL.COM")
kv(doc, "Mob. No. :-  ",      "____________________")
kv(doc, "Agency Code :-  ",   "____________________", sa=14)

para(doc, "संलग्न :– उपरोक्तानुसार।", size=12, sa=28)

# ── Signature ─────────────────────────────────────────────────────────────────
para(doc, "ई–हस्ताक्षर",
     align=WD_ALIGN_PARAGRAPH.RIGHT, size=12, sa=20)
para(doc, "(अनिल खिची)",
     bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, size=12, sa=2)
para(doc, "अधिशासी अभियन्ता",
     align=WD_ALIGN_PARAGRAPH.RIGHT, size=12, sa=2)
para(doc, "सा.नि.वि. जिला खण्ड द्वितीय, उदयपुर",
     align=WD_ALIGN_PARAGRAPH.RIGHT, size=12, sa=0)

doc.save(DOCX_PATH)
print(f"DOCX saved : {DOCX_PATH}")

# ── PDF via Word COM ──────────────────────────────────────────────────────────
try:
    import win32com.client
    tmp  = tempfile.mkdtemp()
    tdx  = os.path.join(tmp, "sna_temp.docx")
    tpdf = os.path.join(tmp, "sna_out.pdf")
    shutil.copy2(DOCX_PATH, tdx)
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    wd = word.Documents.Open(tdx)
    wd.ExportAsFixedFormat(
        OutputFileName=tpdf, ExportFormat=17,
        OpenAfterExport=False, OptimizeFor=0, Range=0,
        IncludeDocProps=True, KeepIRM=True, CreateBookmarks=0,
        DocStructureTags=True, BitmapMissingFonts=True, UseISO19005_1=False)
    wd.Close(False)
    word.Quit()
    shutil.copy2(tpdf, PDF_PATH)
    print(f"PDF  saved : {PDF_PATH}  ({os.path.getsize(PDF_PATH):,} bytes)")
except Exception as e:
    print(f"PDF ERROR  : {e}")

# ── Publish to public/docs/ + show Netlify URL ───────────────────────────────
# SNA letter has no HTML yet — publish DOCX + PDF only.
# To add HTML: generate one and pass html_path= here.
publish_output(docx_path=DOCX_PATH, pdf_path=PDF_PATH)
